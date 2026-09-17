import json
import os
import random
import re
import shutil
import sqlite3
import sys
import tempfile
import threading
import time
import urllib.request
import urllib.parse
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi import FastAPI, UploadFile, File, Form, Query, HTTPException, Body, Request
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from starlette.staticfiles import StaticFiles as StarletteStaticFiles
from pydantic import BaseModel

from backend import db, ai, tts, importer, phrasal, updater, net

db.init_db()

APP_VERSION = '0.2.0c'
GITHUB_REPO = 'HoweyYang/KTRT'
FRONTEND = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'frontend', 'static')
updater.configure(APP_VERSION, sys.executable, bool(getattr(sys, 'frozen', False)))
app = FastAPI(title='溯源词斩 KTRT')

# 本地应用只服务同源页面：原先 CORS 通配等于把本地接口对任意网站敞开，
# 任何网页都能读到设置里的 API Key，也能删词书。去掉通配后浏览器默认不允许
# 跨站读取；这里再挡一道来源检查，把 form / no-cors 这类不受 CORS 约束的
# 跨站写入也一并拦掉。
_LOCAL_ORIGIN = re.compile(r'^http://(?:127\.0\.0\.1|localhost|\[::1\])(?::\d+)?$')


@app.middleware('http')
async def _block_foreign_origin(request: Request, call_next):
    origin = request.headers.get('origin')
    if origin and not _LOCAL_ORIGIN.match(origin):
        return JSONResponse({'detail': '拒绝跨站请求'}, status_code=403)
    return await call_next(request)


class NoCacheStaticFiles(StarletteStaticFiles):
    async def get_response(self, path: str, scope):
        resp = await super().get_response(path, scope)
        resp.headers['Cache-Control'] = 'no-store, must-revalidate'
        return resp

    def lookup_path(self, path):
        """先看数据目录里的热补丁（更新页下载的 patch），再回落到内置前端。"""
        if path.startswith(('/', '\\')):
            return '', None
        bases = [updater.overlay_dir()]
        for d in (getattr(self, 'all_directories', None) or [getattr(self, 'directory', '')]):
            bases.append(d)
        for base in bases:
            if not base or not os.path.isdir(base):
                continue
            root = os.path.realpath(base)
            full = os.path.realpath(os.path.join(root, path))
            try:
                if os.path.commonpath([full, root]) != root:
                    continue
            except ValueError:
                continue
            try:
                return full, os.stat(full)
            except (FileNotFoundError, NotADirectoryError):
                continue
        return '', None


app.mount('/static', NoCacheStaticFiles(directory=FRONTEND), name='static')

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _data_file(name):
    """数据文件定位：用户数据目录 → 源码 data/ → 打包解压目录。"""
    cands = [os.path.join(db.DATA_DIR, name), os.path.join(PROJECT_ROOT, 'data', name)]
    meipass = getattr(sys, '_MEIPASS', '')
    if meipass:
        cands.append(os.path.join(meipass, 'data', name))
    for p in cands:
        if p and os.path.exists(p):
            return p
    return ''


# ---------- 内置动词短语库（data/phrasal_verbs.json，由 tools/enrich_phrasal_verbs.py 生成） ----------
phrasal.load(_data_file('phrasal_verbs.json'))


def _backfill_phrasal_keys():
    """老词库补检索：按动词给每条写入命中的短语动词（只跑一次）。"""
    if not phrasal.BY_HEAD:
        return
    try:
        if db.get_setting('phrasal_backfill') == '1':
            return
        with db._lock:
            with db.get_conn() as conn:
                hit = 0
                for r in conn.execute('SELECT id, word, meaning FROM words'):
                    keys = phrasal.keys_for(r['word'], r['meaning'])
                    if keys:
                        conn.execute('UPDATE words SET phrasal_keys=? WHERE id=?',
                                     (';'.join(keys), r['id']))
                        hit += 1
                conn.execute(
                    "INSERT INTO settings(key, value) VALUES('phrasal_backfill','1') "
                    "ON CONFLICT(key) DO UPDATE SET value='1'")
                if hit:
                    print('[KTRT] 动词短语检索完成：%d 条词条命中' % hit)
    except Exception as e:
        print('[KTRT] 动词短语检索失败：%s' % e)


_backfill_phrasal_keys()


def _lists_meta(book_id):
    with db.get_conn() as conn:
        rows = conn.execute(
            'SELECT w.list_no, COUNT(*) total, '
            '(SELECT COUNT(*) FROM word_status s JOIN words w2 ON s.word_id=w2.id '
            ' WHERE w2.book_id=? AND w2.list_no=w.list_no AND s.learned=1) learned '
            'FROM words w WHERE w.book_id=? GROUP BY w.list_no ORDER BY w.list_no',
            (book_id, book_id),
        ).fetchall()
    return [{'list_no': r['list_no'], 'total': r['total'], 'learned': r['learned']} for r in rows]


def _status(conn, word_id):
    row = conn.execute('SELECT * FROM word_status WHERE word_id=?', (word_id,)).fetchone()
    if row is None:
        return {'familiar': 0, 'unfamiliar': 0, 'favorite': 0, 'learned': 0}
    return {k: row[k] for k in ('familiar', 'unfamiliar', 'favorite', 'learned')}


@app.get('/')
def index():
    root = updater.overlay_dir() if updater.overlay_active() else FRONTEND
    resp = FileResponse(os.path.join(root, 'index.html'))
    resp.headers['Cache-Control'] = 'no-store, must-revalidate'
    return resp


@app.get('/api/ping')
def ping():
    """启动器探针：确认这个端口上跑的确实是 KTRT，而不是别的程序。"""
    return {'app': 'KTRT', 'version': APP_VERSION}


@app.get('/api/bootstrap')
def bootstrap():
    with db.get_conn() as conn:
        books = conn.execute(
            'SELECT b.*, (SELECT COUNT(*) FROM words w WHERE w.book_id=b.id) word_count '
            'FROM word_books b ORDER BY b.id'
        ).fetchall()
    return {
        'books': [dict(r) for r in books],
        'settings': _get_settings(),
        'presets': ai.PRESETS,
    }


@app.get('/api/books/{book_id}/lists')
def book_lists(book_id: int):
    return _lists_meta(book_id)


@app.get('/api/card')
def card(book_id: int = Query(...), list_no: int = Query(1), seq: int = Query(1)):
    with db.get_conn() as conn:
        w = conn.execute(
            'SELECT * FROM words WHERE book_id=? AND list_no=? AND seq=?',
            (book_id, list_no, seq),
        ).fetchone()
        if w is None:
            raise HTTPException(404, '未找到该词')
        status = _status(conn, w['id'])
        sents = conn.execute(
            'SELECT * FROM sentences WHERE word_id=? ORDER BY id', (w['id'],)
        ).fetchall()
        count = conn.execute(
            'SELECT COUNT(*) c FROM words WHERE book_id=? AND list_no=?',
            (book_id, list_no),
        ).fetchone()['c']
        learned = conn.execute(
            'SELECT COUNT(*) c FROM word_status s JOIN words w2 ON s.word_id=w2.id '
            'WHERE w2.book_id=? AND w2.list_no=? AND s.learned=1',
            (book_id, list_no),
        ).fetchone()['c']
    return {
        'word': dict(w),
        'status': status,
        'sentences': [dict(s) for s in sents],
        'progress': {'total': count, 'learned': learned},
    }


class StatusBody(BaseModel):
    word_id: int
    field: str  # learned | familiar | unfamiliar | favorite
    value: bool


@app.post('/api/status')
def set_status(body: StatusBody):
    if body.field not in ('learned', 'familiar', 'unfamiliar', 'favorite'):
        raise HTTPException(400, '未知状态字段')
    with db._lock:
        with db.get_conn() as conn:
            conn.execute(
                'INSERT INTO word_status(word_id, familiar, unfamiliar, favorite, learned, updated_at) '
                'VALUES(?,0,0,0,0,datetime(\'now\',\'localtime\')) '
                'ON CONFLICT(word_id) DO NOTHING',
                (body.word_id,),
            )
            conn.execute(
                f'UPDATE word_status SET {body.field}=?, updated_at=datetime(\'now\',\'localtime\') '
                'WHERE word_id=?',
                (1 if body.value else 0, body.word_id),
            )
            # 已背与不熟悉互斥：点亮其一，自动熄灭另一个
            if body.field == 'learned' and body.value:
                conn.execute('UPDATE word_status SET unfamiliar=0 WHERE word_id=?', (body.word_id,))
            elif body.field == 'unfamiliar' and body.value:
                conn.execute('UPDATE word_status SET learned=0 WHERE word_id=?', (body.word_id,))
            status = _status(conn, body.word_id)
    return {'status': status}


@app.delete('/api/status/{word_id}')
def delete_status(word_id: int):
    with db._lock:
        with db.get_conn() as conn:
            conn.execute('DELETE FROM word_status WHERE word_id=?', (word_id,))
    return {'ok': True}


class ClearBody(BaseModel):
    book_id: int
    list_no: int


@app.post('/api/lists/clear')
def clear_list(body: ClearBody):
    with db._lock:
        with db.get_conn() as conn:
            conn.execute(
                'UPDATE word_status SET learned=0 WHERE word_id IN '
                '(SELECT id FROM words WHERE book_id=? AND list_no=?)',
                (body.book_id, body.list_no),
            )
    return {'ok': True}


class SentenceBody(BaseModel):
    word_id: int
    prompt: str = ''


def _extract_json(text):
    text = re.sub(r'```(?:json)?', '', text).strip()
    a, b = text.find('{'), text.rfind('}')
    if a == -1 or b == -1:
        raise ValueError('AI 未返回 JSON')
    return json.loads(text[a:b + 1])


@app.post('/api/sentences')
def create_sentence(body: SentenceBody):
    with db.get_conn() as conn:
        w = conn.execute('SELECT * FROM words WHERE id=?', (body.word_id,)).fetchone()
        if w is None:
            raise HTTPException(404, '未找到该词')
    prompt = (body.prompt or '').strip()
    messages = [
        {'role': 'system', 'content': '你是英语造句老师。只输出 JSON，不要任何多余文字。'},
        {'role': 'user', 'content': (
            f'用英文为单词「{w["word"]}」造句，规则：'
            '1) 先判断中文提示词是否与该单词语义相关；'
            '2) 若提示词与单词无关，完全忽略提示词，基于该单词随机造一个自然句子；'
            '3) 若提示词相关，按提示词的意思造句，句中可以使用该单词，也可以使用其同根词或不同词性变形'
            f'（如动词变名词、名词变形容词等，示例：abbreviate → abbreviation），只要读者能看出与「{w["word"]}」相关即可；'
            '4) 句子自然地道、长度适中。'
            f'中文提示词：{prompt or "（无）"}。'
            '输出 JSON：{"sentence": "英文句子", "translation": "整句中文翻译"}'
        )},
    ]
    try:
        raw = ai.chat(messages, max_tokens=500, temperature=0.7)
        data = _extract_json(raw)
        sentence = data.get('sentence', '').strip()
        translation = data.get('translation', '').strip()
    except Exception as e:
        raise HTTPException(502, f'AI 生成失败：{e}')
    if not sentence:
        raise HTTPException(502, 'AI 返回了空句子')
    with db._lock:
        with db.get_conn() as conn:
            conn.execute(
                'INSERT INTO sentences(word_id, prompt, sentence, translation, created_at) '
                'VALUES(?,?,?,?,datetime(\'now\',\'localtime\'))',
                (body.word_id, prompt, sentence, translation),
            )
            # 每词最多 3 句，超出删除最旧
            conn.execute(
                'DELETE FROM sentences WHERE id NOT IN '
                '(SELECT id FROM sentences WHERE word_id=? ORDER BY id DESC LIMIT 3) '
                'AND word_id=?',
                (body.word_id, body.word_id),
            )
            sents = conn.execute(
                'SELECT * FROM sentences WHERE word_id=? ORDER BY id', (body.word_id,)
            ).fetchall()
    return {'sentences': [dict(s) for s in sents]}


@app.delete('/api/sentences/{sentence_id}')
def delete_sentence(sentence_id: int):
    with db._lock:
        with db.get_conn() as conn:
            conn.execute('DELETE FROM sentences WHERE id=?', (sentence_id,))
    return {'ok': True}


@app.get('/api/sentences')
def list_sentences(word_id: int = Query(...)):
    with db.get_conn() as conn:
        sents = conn.execute(
            'SELECT * FROM sentences WHERE word_id=? ORDER BY id', (word_id,)
        ).fetchall()
    return [dict(s) for s in sents]


class NoteBody(BaseModel):
    word_id: int
    content: str = ''


class NotesClearBody(BaseModel):
    book_id: int = 0
    list_no: int = 0


@app.get('/api/notes/{word_id}')
def get_note(word_id: int):
    with db.get_conn() as conn:
        w = conn.execute('SELECT word FROM words WHERE id=?', (word_id,)).fetchone()
        key = _note_key(w['word']) if w else ''
        row = conn.execute(
            'SELECT content, updated_at FROM word_notes WHERE word_key=?', (key,)
        ).fetchone() if key else None
    return {'word_id': word_id, 'word_key': key,
            'content': row['content'] if row else '',
            'updated_at': row['updated_at'] if row else ''}


@app.post('/api/notes')
def save_note(body: NoteBody):
    content = (body.content or '').strip()
    with db._lock:
        with db.get_conn() as conn:
            w = conn.execute('SELECT word FROM words WHERE id=?', (body.word_id,)).fetchone()
            if w is None:
                raise HTTPException(404, '未找到该词')
            key = _note_key(w['word'])
            if content:
                conn.execute(
                    'INSERT INTO word_notes(word_key, content, updated_at) '
                    'VALUES(?,?,datetime(\'now\',\'localtime\')) '
                    'ON CONFLICT(word_key) DO UPDATE SET content=excluded.content, '
                    'updated_at=excluded.updated_at',
                    (key, content),
                )
            else:
                conn.execute('DELETE FROM word_notes WHERE word_key=?', (key,))
    return {'ok': True, 'content': content, 'word_key': key}


@app.delete('/api/notes/{word_id}')
def delete_note(word_id: int):
    with db._lock:
        with db.get_conn() as conn:
            w = conn.execute('SELECT word FROM words WHERE id=?', (word_id,)).fetchone()
            if w is not None:
                conn.execute('DELETE FROM word_notes WHERE word_key=?', (_note_key(w['word']),))
    return {'ok': True}


@app.post('/api/notes/clear')
def clear_notes(body: NotesClearBody):
    sql = 'SELECT word FROM words w WHERE 1=1'
    params = []
    if body.book_id:
        sql += ' AND w.book_id=?'
        params.append(body.book_id)
    if body.list_no:
        sql += ' AND w.list_no=?'
        params.append(body.list_no)
    with db._lock:
        with db.get_conn() as conn:
            keys = {_note_key(r['word']) for r in conn.execute(sql, params)}
            keys.discard('')
            cleared = 0
            for k in keys:
                cur = conn.execute('DELETE FROM word_notes WHERE word_key=?', (k,))
                cleared += cur.rowcount or 0
    return {'ok': True, 'cleared': cleared}


class BookmarkBody(BaseModel):
    book_id: int
    list_no: int
    seq: int
    word: str = ''


@app.get('/api/bookmarks')
def list_bookmarks():
    with db.get_conn() as conn:
        rows = conn.execute(
            'SELECT bm.id, bm.book_id, b.name AS book_name, bm.list_no, bm.seq, '
            'bm.word, bm.created_at FROM bookmarks bm '
            'JOIN word_books b ON b.id = bm.book_id '
            'ORDER BY bm.book_id, bm.list_no, bm.seq'
        ).fetchall()
    return [dict(r) for r in rows]


@app.post('/api/bookmarks')
def add_bookmark(body: BookmarkBody):
    with db._lock:
        with db.get_conn() as conn:
            row = conn.execute(
                'SELECT word FROM words WHERE book_id=? AND list_no=? AND seq=?',
                (body.book_id, body.list_no, body.seq),
            ).fetchone()
            if row is None:
                raise HTTPException(404, '该位置已不存在，无法添加书签')
            conn.execute(
                'INSERT INTO bookmarks(book_id, list_no, seq, word) VALUES(?,?,?,?) '
                'ON CONFLICT(book_id, list_no, seq) DO UPDATE SET word=excluded.word',
                (body.book_id, body.list_no, body.seq, row['word']),
            )
            rec = conn.execute(
                'SELECT id FROM bookmarks WHERE book_id=? AND list_no=? AND seq=?',
                (body.book_id, body.list_no, body.seq),
            ).fetchone()
    return {'ok': True, 'id': rec['id']}


@app.delete('/api/bookmarks/{bookmark_id}')
def delete_bookmark(bookmark_id: int):
    with db._lock:
        with db.get_conn() as conn:
            conn.execute('DELETE FROM bookmarks WHERE id=?', (bookmark_id,))
    return {'ok': True}


def _challenge_options(conn, book_id, word_ids):
    """给一组 word_id 生成四选一题目：1 正确 + 3 错误（错误项来自本书其他词的中文释义）。"""
    meanings = [r['meaning'] for r in conn.execute(
        'SELECT meaning FROM words WHERE book_id=? AND TRIM(COALESCE(meaning,""))<>""', (book_id,)).fetchall()]
    questions = []
    for wid in word_ids:
        w = conn.execute('SELECT word, meaning FROM words WHERE id=?', (wid,)).fetchone()
        if w is None or not (w['meaning'] or '').strip():
            continue
        correct = w['meaning'].strip()
        pool = [m.strip() for m in meanings if m.strip() != correct]
        wrong = random.sample(pool, min(3, len(pool)))
        opts = [correct] + wrong
        random.shuffle(opts)
        questions.append({
            'word_id': wid,
            'word': w['word'],
            'options': opts,
            'correct': opts.index(correct),
        })
    random.shuffle(questions)
    return questions


@app.get('/api/challenge')
def challenge(book_id: int = Query(...), list_no: int = Query(...)):
    with db.get_conn() as conn:
        ids = [r['id'] for r in conn.execute(
            'SELECT w.id FROM words w LEFT JOIN word_status s ON s.word_id=w.id '
            'WHERE w.book_id=? AND w.list_no=? AND COALESCE(s.learned,0)=0 ORDER BY w.seq',
            (book_id, list_no)).fetchall()]
        questions = _challenge_options(conn, book_id, ids)
        best = conn.execute('SELECT best FROM challenge_scores WHERE book_id=? AND list_no=?',
                            (book_id, list_no)).fetchone()
    return {'questions': questions, 'total': len(questions), 'best': best['best'] if best else 0, 'list_total': len(ids)}


class ChallengeResultBody(BaseModel):
    book_id: int
    list_no: int
    wrong_ids: list = []
    total: int = 0
    correct: int = 0


@app.post('/api/challenge/result')
def challenge_result(body: ChallengeResultBody):
    """记词闯关结算：只更新最高分（错词已由 /api/mistakes/add 即时入库）。"""
    with db._lock:
        with db.get_conn() as conn:
            conn.execute(
                'INSERT INTO challenge_scores(book_id, list_no, best) VALUES(?,?,?) '
                'ON CONFLICT(book_id, list_no) DO UPDATE SET best=MAX(best, excluded.best), '
                'updated_at=datetime(\'now\',\'localtime\')',
                (body.book_id, body.list_no, body.correct),
            )
    with db.get_conn() as conn:
        best = conn.execute('SELECT best FROM challenge_scores WHERE book_id=? AND list_no=?',
                            (body.book_id, body.list_no)).fetchone()
    return {'ok': True, 'best': best['best'] if best else 0}


class MistakesAddBody(BaseModel):
    word_id: int


@app.post('/api/mistakes/add')
def mistake_add(body: MistakesAddBody):
    """记词闯关答错时即时调用：标不熟悉 + 入错题本（切换/退出不影响已入册的错词）。"""
    with db._lock:
        with db.get_conn() as conn:
            w = conn.execute('SELECT book_id, list_no FROM words WHERE id=?', (body.word_id,)).fetchone()
            if w is None:
                raise HTTPException(404, '未找到该词')
            conn.execute('INSERT OR IGNORE INTO word_status(word_id) VALUES(?)', (body.word_id,))
            conn.execute('UPDATE word_status SET unfamiliar=1, learned=0 WHERE word_id=?', (body.word_id,))
            conn.execute('INSERT OR IGNORE INTO mistakes(word_id, book_id, list_no) VALUES(?,?,?)',
                         (body.word_id, w['book_id'], w['list_no']))
    return {'ok': True}


@app.get('/api/mistakes')
def mistakes(book_id: int = Query(...)):
    with db.get_conn() as conn:
        rows = conn.execute(
            'SELECT m.word_id, m.list_no, w.word, w.meaning FROM mistakes m '
            'JOIN words w ON w.id=m.word_id LEFT JOIN word_status s ON s.word_id=m.word_id '
            'WHERE m.book_id=? AND COALESCE(s.learned,0)=0 ORDER BY m.list_no, w.seq', (book_id,)).fetchall()
    return [dict(r) for r in rows]


@app.get('/api/mistakes/lists')
def mistake_lists(book_id: int = Query(...)):
    with db.get_conn() as conn:
        rows = conn.execute(
            'SELECT m.list_no, COUNT(*) c FROM mistakes m LEFT JOIN word_status s ON s.word_id=m.word_id '
            'WHERE m.book_id=? AND COALESCE(s.learned,0)=0 GROUP BY m.list_no ORDER BY m.list_no', (book_id,)).fetchall()
    return [dict(r) for r in rows]


class MistakesStartBody(BaseModel):
    book_id: int
    list_nos: list = []


@app.post('/api/mistakes/start')
def mistakes_start(body: MistakesStartBody):
    ids = []
    if body.list_nos:
        ph = ','.join('?' * len(body.list_nos))
        with db.get_conn() as conn:
            ids = [r['word_id'] for r in conn.execute(
                'SELECT m.word_id FROM mistakes m LEFT JOIN word_status s ON s.word_id=m.word_id '
                'WHERE m.book_id=? AND m.list_no IN (%s) AND COALESCE(s.learned,0)=0' % ph,
                [body.book_id] + body.list_nos).fetchall()]
    with db.get_conn() as conn:
        questions = _challenge_options(conn, body.book_id, ids)
    return {'questions': questions, 'total': len(questions)}


class MistakesResultBody(BaseModel):
    removed_ids: list = []


@app.post('/api/mistakes/result')
def mistakes_result(body: MistakesResultBody):
    with db._lock:
        with db.get_conn() as conn:
            for wid in body.removed_ids:
                conn.execute('DELETE FROM mistakes WHERE word_id=?', (wid,))
    return {'ok': True}


@app.get('/api/manage')
def manage(filter: str = Query('all')):
    where = ''
    params = []
    if filter == 'learned':
        where = 'WHERE s.learned=1'
    elif filter == 'familiar':
        where = 'WHERE s.familiar=1'
    elif filter == 'unfamiliar':
        where = 'WHERE s.unfamiliar=1'
    elif filter == 'favorite':
        where = 'WHERE s.favorite=1'
    elif filter == 'sentences':
        where = 'WHERE (SELECT COUNT(*) FROM sentences x WHERE x.word_id=w.id)>0'
    elif filter == 'notes':
        pass  # 笔记按词共享，下面在 Python 侧过滤
    with db.get_conn() as conn:
        note_keys = {r['word_key'] for r in conn.execute('SELECT word_key FROM word_notes')}
        rows = conn.execute(
            'SELECT w.id, w.word, w.phonetic, w.list_no, w.seq, b.name book_name, b.language, '
            's.familiar, s.unfamiliar, s.favorite, s.learned, '
            '(SELECT COUNT(*) FROM sentences x WHERE x.word_id=w.id) sent_count '
            f'FROM words w JOIN word_books b ON b.id=w.book_id '
            f'LEFT JOIN word_status s ON s.word_id=w.id {where} ORDER BY b.id, w.list_no, w.seq',
            params,
        ).fetchall()
    out = []
    for r in rows:
        d = dict(r)
        d['has_note'] = 1 if _note_key(d['word']) in note_keys else 0
        if filter == 'notes' and not d['has_note']:
            continue
        out.append(d)
    return out


@app.get('/api/export')
def export_words(scope: str = Query('unfamiliar'), book_id: int = Query(None), list_no: int = Query(None)):
    if scope not in ('unfamiliar', 'favorite', 'both'):
        raise HTTPException(400, '未知导出范围')
    conds, params = [], []
    if scope == 'favorite':
        conds.append('s.favorite=1')
    elif scope == 'both':
        conds.append('(s.unfamiliar=1 OR s.favorite=1)')
    else:
        conds.append('s.unfamiliar=1')
    if book_id is not None:
        conds.append('w.book_id=?')
        params.append(book_id)
    if list_no is not None:
        conds.append('w.list_no=?')
        params.append(list_no)
    where = 'WHERE ' + ' AND '.join(conds)
    with db.get_conn() as conn:
        rows = conn.execute(
            'SELECT w.word, w.phonetic, w.meaning, w.collocations, w.phrases, '
            'w.synonyms, w.antonyms, w.root_words, w.list_no, b.name AS book_name, b.language '
            'FROM words w JOIN word_books b ON b.id=w.book_id '
            f'LEFT JOIN word_status s ON s.word_id=w.id {where} ORDER BY b.id, w.list_no, w.seq',
            params,
        ).fetchall()
    if not rows:
        raise HTTPException(404, '没有符合条件可导出的词汇')
    from openpyxl import Workbook
    from datetime import datetime
    wb = Workbook()
    ws = wb.active
    ws.title = '词汇'
    ws.append(['【单词】', '【音标】', '【词性释义】', '【搭配】', '【短语】', '【同义词】', '【反义词】', '【同根词】', '【List】', '【语言】', '【单词书】'])
    for r in rows:
        ws.append([r['word'], r['phonetic'], r['meaning'], r['collocations'], r['phrases'], r['synonyms'], r['antonyms'], r['root_words'], r['list_no'], r['language'], r['book_name']])
    label = {'unfamiliar': '不熟悉', 'favorite': '收藏', 'both': '不熟悉与收藏'}[scope]
    name = f"KTRT_{label}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = os.path.join(tempfile.gettempdir(), name)
    wb.save(path)
    return FileResponse(path, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename=name)


@app.get('/api/export/notes')
def export_notes(book_id: int = Query(None), list_no: int = Query(None)):
    """导出笔记：只含写了笔记的词，Markdown 文件，格式为「词 + 笔记」循环。"""
    conds, params = [], []
    if book_id is not None:
        conds.append('w.book_id=?')
        params.append(book_id)
    if list_no is not None:
        conds.append('w.list_no=?')
        params.append(list_no)
    where = ('WHERE ' + ' AND '.join(conds)) if conds else ''
    with db.get_conn() as conn:
        notes = {r['word_key']: (r['content'] or '')
                 for r in conn.execute('SELECT word_key, content FROM word_notes')}
        rows = conn.execute(
            'SELECT w.word FROM words w '
            'JOIN word_books b ON b.id=w.book_id '
            f'{where} ORDER BY b.id, w.list_no, w.seq',
            params,
        ).fetchall()
    from datetime import datetime
    parts, seen = [], set()
    for r in rows:
        key = _note_key(r['word'])
        if not key or key in seen or key not in notes:
            continue
        seen.add(key)
        parts.append('## ' + r['word'] + '\n\n' + notes[key] + '\n')
    if not parts:
        raise HTTPException(404, '没有可导出的笔记')
    text = '\n'.join(parts)
    name = f"KTRT_笔记_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md"
    path = os.path.join(tempfile.gettempdir(), name)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(text)
    return FileResponse(path, media_type='text/markdown', filename=name)


# ---------- 词书本地副本与词条编辑 ----------

EDIT_FIELDS = ('word', 'phonetic', 'meaning', 'collocations', 'phrases',
               'synonyms', 'antonyms', 'root_words')
EXCEL_HEADERS = {
    'word': '【单词】', 'phonetic': '【音标】', 'meaning': '【词性释义】',
    'collocations': '【搭配】', 'phrases': '【短语】', 'synonyms': '【同义词】',
    'antonyms': '【反义词】', 'root_words': '【同根词】',
}
BOOK_FILES_DIR = os.path.join(db.DATA_DIR, 'wordbooks')


def _safe_filename(name):
    return re.sub(r'[\\/:*?"<>|]', '_', (name or '').strip()) or '词书'


def _keep_uploaded_book(src, book_name, suffix):
    """导入时留一份词书到本地词书目录，并把它记为这本书的来源文件。"""
    os.makedirs(BOOK_FILES_DIR, exist_ok=True)
    dest = os.path.join(BOOK_FILES_DIR, _safe_filename(book_name) + suffix)
    shutil.copyfile(src, dest)
    with db._lock:
        with db.get_conn() as conn:
            conn.execute('UPDATE word_books SET source=? WHERE name=?', (dest, book_name))
    return dest


def _export_book_xlsx(conn, book, path):
    """按 11 列【】格式把整本书从本地库导出成 xlsx（原文件失效时的兜底副本）。"""
    from openpyxl import Workbook
    rows = conn.execute(
        'SELECT word, phonetic, meaning, collocations, phrases, synonyms, antonyms, '
        'root_words, list_no FROM words WHERE book_id=? ORDER BY list_no, seq',
        (book['id'],),
    ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = _safe_filename(book['name'])[:31]
    ws.append(['【单词】', '【音标】', '【词性释义】', '【搭配】', '【短语】', '【同义词】',
               '【反义词】', '【同根词】', '【List】', '【语言】', '【单词书】'])
    for r in rows:
        ws.append([r['word'], r['phonetic'], r['meaning'], r['collocations'], r['phrases'],
                   r['synonyms'], r['antonyms'], r['root_words'], r['list_no'],
                   book['language'], book['name']])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    conn.execute('UPDATE word_books SET source=? WHERE id=?', (path, book['id']))
    return path


def _book_excel_path(conn, book):
    """返回该词书可写的 Excel 路径：优先原文件，失效则落本地托管副本。"""
    src = (book['source'] or '').strip()
    if src and os.path.isfile(src):
        return src, False
    path = os.path.join(BOOK_FILES_DIR, _safe_filename(book['name']) + '.xlsx')
    _export_book_xlsx(conn, book, path)
    return path, True


def _write_excel_entry(path, old_word, list_no, changed):
    """把改动写回词书 Excel 的对应行，返回 (是否写入, 说明)。"""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active

    header_row, colmap = 1, {}
    for r in range(1, min(5, ws.max_row) + 1):
        cand = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            h = str(v).strip()
            for field, head in EXCEL_HEADERS.items():
                if h == head:
                    cand[field] = c
        if 'word' in cand:
            header_row, colmap = r, cand
            break
    if 'word' not in colmap:
        wb.close()
        return False, '词书缺少【单词】列，已只改本地库'

    list_col = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v is not None and str(v).strip() in ('【List】', 'List'):
            list_col = c
            break

    target = None
    old = (old_word or '').strip().lower()
    for r in range(header_row + 1, ws.max_row + 1):
        wv = ws.cell(r, colmap['word']).value
        if wv is None or str(wv).strip().lower() != old:
            continue
        if list_col is not None and list_no is not None:
            lv = ws.cell(r, list_col).value
            if lv is not None and str(lv).strip() not in ('', str(list_no)):
                continue
        target = r
        break
    if target is None:
        wb.close()
        return False, '词书 Excel 中未找到「%s」，已只改本地库' % old_word

    for field, val in changed.items():
        if field in colmap:
            ws.cell(target, colmap[field]).value = val
    wb.save(path)
    wb.close()
    return True, ''


@app.post('/api/import')
async def import_book(
    file: UploadFile = File(...),
    book_name: str = Form(''),
    language: str = Form('英语'),
):
    suffix = os.path.splitext(file.filename or '')[-1] or '.xlsx'
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(await file.read())
        path = tmp.name
    try:
        result = importer.import_book(path, book_name.strip(), language.strip())
        try:
            _keep_uploaded_book(path, result['book_name'], suffix)
        except Exception:
            pass  # 留副本失败不影响导入本身
        return result
    except Exception as e:
        raise HTTPException(400, str(e))
    finally:
        try:
            os.unlink(path)
        except Exception:
            pass


class EditBody(BaseModel):
    word: str | None = None
    phonetic: str | None = None
    meaning: str | None = None
    collocations: str | None = None
    phrases: str | None = None
    synonyms: str | None = None
    antonyms: str | None = None
    root_words: str | None = None


@app.post('/api/word/{word_id}/edit')
def edit_word(word_id: int, body: EditBody):
    """编辑词条：同步本地库，并回写这个词书 Excel 的对应行。"""
    changed = {}
    for k in EDIT_FIELDS:
        v = getattr(body, k)
        if v is not None:
            changed[k] = v.strip()
    if not changed:
        raise HTTPException(400, '没有要保存的内容')
    if 'word' in changed and not changed['word']:
        raise HTTPException(400, '单词不能为空')

    with db._lock:
        with db.get_conn() as conn:
            w = conn.execute('SELECT * FROM words WHERE id=?', (word_id,)).fetchone()
            if w is None:
                raise HTTPException(404, '词条不存在')
            book = conn.execute('SELECT * FROM word_books WHERE id=?', (w['book_id'],)).fetchone()
            conn.execute(
                'UPDATE words SET ' + ', '.join(f'{k}=?' for k in changed) + ' WHERE id=?',
                (*changed.values(), word_id),
            )
            excel = {'updated': False, 'created': False, 'path': '', 'message': ''}
            if book is not None and book['name'] != '外部单词收藏册':
                try:
                    path, created = _book_excel_path(conn, book)
                    excel['path'] = path
                    excel['created'] = created
                    ok, msg = _write_excel_entry(path, w['word'], w['list_no'], changed)
                    excel['updated'] = ok
                    excel['message'] = msg
                except PermissionError:
                    excel['message'] = '词书 Excel 正被占用（若已在 Excel 里打开请先关闭），已只改本地库'
                except Exception as e:
                    excel['message'] = '回写词书失败：' + str(e)
    return {'ok': True, 'word_id': word_id, 'changed': sorted(changed), 'excel': excel}


@app.delete('/api/books/{book_id}')
def delete_book(book_id: int):
    with db._lock:
        with db.get_conn() as conn:
            book = conn.execute('SELECT * FROM word_books WHERE id=?', (book_id,)).fetchone()
            if book is None:
                raise HTTPException(404, '单词书不存在')
            if book['name'] == '外部单词收藏册':
                raise HTTPException(400, '外部单词收藏册 为默认单词书，不可删除')
            conn.execute('DELETE FROM word_books WHERE id=?', (book_id,))
    return {'ok': True, 'deleted': book['name']}


@app.get('/api/dict/{word}')
def lookup(word: str):
    import sqlite3
    if not os.path.exists(db.DICT_DB_PATH):
        return {'available': False, 'message': '离线词典未导入'}
    try:
        conn = sqlite3.connect(db.DICT_DB_PATH)
        row = conn.execute(
            'SELECT * FROM dict WHERE word=? COLLATE NOCASE', (word.lower(),)
        ).fetchone()
        conn.close()
    except Exception:
        return {'available': False, 'message': '词典读取失败'}
    if row is None:
        return {'available': True, 'found': False}
    return {
        'available': True,
        'found': True,
        'word': row[0],
        'phonetic': row[1] or '',
        'definition': row[2] or '',
        'translation': row[3] or '',
        'pos': row[4] or '',
        'exchange': row[5] or '',
        'exchange_text': _exchange_pretty(row[5] or ''),
        'collins': row[6] or '',
        'oxford': row[7] or '',
        'tag': row[8] or '',
        'bnc': row[9] or '',
        'frq': row[10] or '',
    }


class CustomDictBody(BaseModel):
    word: str = ''


def _morph_candidates(w):
    """规则兜底：由表面词形生成可能的原形候选（ECDICT 0: 覆盖不到时用）。"""
    cands = []
    n = len(w)
    if w.endswith('ies') and n > 4:
        cands += [w[:-3] + 'y', w[:-1]]
    if w.endswith('es') and n > 3:
        cands += [w[:-2], w[:-1]]
    elif w.endswith('s') and not w.endswith('ss') and n > 2:
        cands.append(w[:-1])
    if w.endswith('ied') and n > 4:
        cands += [w[:-3] + 'y', w[:-1]]
    if w.endswith('ed') and n > 3:
        cands += [w[:-2], w[:-1]]
        if n > 4 and w[-3] == w[-4]:
            cands.append(w[:-3])
    if w.endswith('ing') and n > 4:
        base = w[:-3]
        cands += [base, base + 'e']
        if len(base) >= 3 and base[-2] == base[-1] and base[-1] not in 'aeiou':
            cands.append(base[:-1])
    return [c for c in dict.fromkeys(cands) if c and re.match(r"^[a-z]+$", c)]


def _canonical_word(raw):
    """把搜索词归一到原形：先看已导入词书，再看 ECDICT 的 0: 原形字段与规则兜底。"""
    raw = (raw or '').strip()
    if not raw:
        return raw
    if ' ' in raw:
        return raw  # 多词短语不做单词变形归一
    w = raw.lower()
    with db.get_conn() as conn:
        if conn.execute('SELECT 1 FROM words WHERE word=? COLLATE NOCASE LIMIT 1', (raw,)).fetchone():
            return raw  # 输入本身就是词书里的词（大小写原样保留）
    if not re.match(r"^[a-z][a-z\-']*$", w):
        return raw
    cands = []
    if os.path.exists(db.DICT_DB_PATH):
        try:
            dconn = sqlite3.connect(db.DICT_DB_PATH)
            row = dconn.execute('SELECT * FROM dict WHERE word=? COLLATE NOCASE', (w,)).fetchone()
            if row:
                for tok in (row[5] or '').split('/'):
                    tok = tok.strip()
                    if tok.startswith('0:'):
                        lemma = tok[2:].strip()
                        if lemma and lemma.lower() != w:
                            cands.append(lemma)
            if row is None:
                # 词典里没有这个词条时，才用规则兜底猜原形；
                # 已经有词条（如 bus）说明它本身就是词，绝不能再截尾（bus → bu）。
                for c in _morph_candidates(w):
                    if dconn.execute('SELECT 1 FROM dict WHERE word=? COLLATE NOCASE', (c,)).fetchone():
                        cands.append(c)
            dconn.close()
        except Exception:
            pass
    seen = set()
    with db.get_conn() as conn:
        for c in cands:
            key = c.lower()
            if key in seen:
                continue
            seen.add(key)
            if conn.execute('SELECT 1 FROM words WHERE word=? COLLATE NOCASE LIMIT 1', (c,)).fetchone():
                return c  # 原形在词书里，直接命中
    return cands[0] if cands else raw


_note_key_cache = {}
_dict_ro_conn = None
_note_key_lock = threading.Lock()


def _dict_exchange(word):
    """查 ECDICT 的 exchange 字段（只读、单连接复用；词表是静态的，结果可缓存）。"""
    global _dict_ro_conn
    with _note_key_lock:
        if _dict_ro_conn is None:
            if not os.path.exists(db.DICT_DB_PATH):
                return ''
            try:
                _dict_ro_conn = sqlite3.connect(db.DICT_DB_PATH, check_same_thread=False)
            except Exception:
                return ''
        try:
            row = _dict_ro_conn.execute('SELECT exchange FROM dict WHERE word=?', (word,)).fetchone()
        except Exception:
            return ''
    return (row[0] or '') if row else ''


def _note_key(word):
    """笔记归属键：单词按原形小写（含大小写、复数/变形归一），短语只统一小写与空白。

    同一个词出现在多本词书里时共用同一份笔记；换书、删书都不影响它。
    """
    w = (word or '').strip()
    if not w:
        return ''
    if ' ' in w:
        return re.sub(r'\s+', ' ', w).lower()
    lw = w.lower()
    if lw in _note_key_cache:
        return _note_key_cache[lw]
    key = lw
    ex = _dict_exchange(lw)
    if ex:
        for tok in ex.split('/'):
            tok = tok.strip()
            if tok.startswith('0:') and tok[2:].strip():
                key = tok[2:].strip().lower()
                break
    _note_key_cache[lw] = key
    return key


def _word_positions(conn, word, limit=8):
    """该词出现在哪些词书里、具体位置（book_id/list_no/seq），供界面点击跳转。"""
    w = (word or '').strip()
    if not w:
        return []
    key = _note_key(w)
    rows = conn.execute(
        'SELECT w.id AS word_id, w.word, w.list_no, w.seq, '
        'b.id AS book_id, b.name AS book_name, b.language '
        'FROM words w JOIN word_books b ON b.id=w.book_id '
        'WHERE lower(w.word)=? OR lower(w.word)=? '
        'ORDER BY b.id, w.list_no, w.seq LIMIT ?',
        (key, w.lower(), limit),
    ).fetchall()
    return [dict(r) for r in rows]


def _migrate_notes_to_keys():
    """一次性迁移：把旧的「按词条存」的笔记搬到「按词存」的 word_notes。"""
    try:
        with db._lock:
            with db.get_conn() as conn:
                if conn.execute('SELECT COUNT(*) c FROM word_notes').fetchone()['c']:
                    return
                old = conn.execute(
                    'SELECT n.word_id, n.content, n.updated_at, w.word FROM notes n '
                    'JOIN words w ON w.id=n.word_id'
                ).fetchall()
                for r in old:
                    key = _note_key(r['word'])
                    if not key:
                        continue
                    conn.execute(
                        'INSERT INTO word_notes(word_key, content, updated_at) VALUES(?,?,?) '
                        'ON CONFLICT(word_key) DO UPDATE SET content=excluded.content, '
                        'updated_at=excluded.updated_at',
                        (key, r['content'] or '', r['updated_at']),
                    )
                if old:
                    print('[KTRT] 笔记已按词归并：%d 条' % len(old))
    except Exception as e:
        print('[KTRT] 笔记迁移失败：%s' % e)


_migrate_notes_to_keys()


def _exchange_pretty(ex):
    """把 ECDICT exchange 代码翻译成自然语言。"""
    labels = {'p': '过去式', 'd': '过去分词', 'i': '现在分词', '3': '第三人称单数',
              'r': '比较级', 't': '最高级', 's': '复数'}
    parts = []
    for tok in (ex or '').split('/'):
        tok = tok.strip()
        if ':' not in tok:
            continue
        k, _, v = tok.partition(':')
        k = k.strip()
        v = v.strip()
        if k in labels and v:
            parts.append(labels[k] + ' ' + v)
        elif k == '0' and v:
            parts.append('原形 ' + v)
    return '；'.join(parts)


def _is_phrase(s):
    return ' ' in (s or '').strip()


def _split_phrase(word):
    toks = [t.strip(".,;:!?\"'") for t in (word or '').split()]
    return [t for t in toks if re.match(r"^[A-Za-z][A-Za-z\-']*$", t)]


def _phrase_components(word):
    """把短语拆成组成词，逐个取离线释义与所在词书，供前端展示 / AI 翻译。"""
    out = []
    dconn = None
    if os.path.exists(db.DICT_DB_PATH):
        dconn = sqlite3.connect(db.DICT_DB_PATH)
    for tok in _split_phrase(word)[:8]:
        row = None
        if dconn is not None:
            try:
                row = dconn.execute('SELECT word, phonetic, translation FROM dict WHERE word=? COLLATE NOCASE', (tok.lower(),)).fetchone()
            except Exception:
                row = None
        with db.get_conn() as conn:
            books = [r['name'] for r in conn.execute(
                'SELECT DISTINCT b.name FROM words w JOIN word_books b ON b.id=w.book_id '
                'WHERE w.word=? COLLATE NOCASE', (tok,))]
        out.append({
            'word': (row[0] if row else tok),
            'phonetic': (row[1] or '') if row else '',
            'translation': ((row[2] or '').replace('\n', '；') if row else ''),
            'in_books': books,
        })
    if dconn is not None:
        dconn.close()
    return out


def _damerau(a, b):
    """带换位的编辑距离（小写 ASCII 词足够快）。"""
    n, m = len(a), len(b)
    if abs(n - m) > 3:
        return 99
    d = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n + 1):
        d[i][0] = i
    for j in range(m + 1):
        d[0][j] = j
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            cost = 0 if a[i - 1] == b[j - 1] else 1
            d[i][j] = min(d[i - 1][j] + 1, d[i][j - 1] + 1, d[i - 1][j - 1] + cost)
            if i > 1 and j > 1 and a[i - 1] == b[j - 2] and a[i - 2] == b[j - 1]:
                d[i][j] = min(d[i][j], d[i - 2][j - 2] + 1)
    return d[n][m]


def _suggestions(word, limit=5):
    """本地“你是不是想搜”：词书词全量 + 词典同长度/同首字母高频词，按编辑距离排序。"""
    w = word.lower()
    scored = []
    seen = set()
    with db.get_conn() as conn:
        for r in conn.execute('SELECT DISTINCT lower(word) w FROM words WHERE word <> ? COLLATE NOCASE', (word,)):
            c = r['w']
            if c and c not in seen and abs(len(c) - len(w)) <= 2:
                seen.add(c)
                d = _damerau(w, c)
                if d <= 2:
                    scored.append((d - .5, c))  # 词书里的词优先
    if os.path.exists(db.DICT_DB_PATH):
        try:
            dconn = sqlite3.connect(db.DICT_DB_PATH)
            first = w[0] if w else 'a'
            rows = dconn.execute(
                'SELECT word, COALESCE(bnc,0) b FROM dict '
                'WHERE word LIKE ? AND length(word) BETWEEN ? AND ? '
                'ORDER BY b DESC LIMIT 2500',
                (first + '%', len(w) - 1, len(w) + 1),
            ).fetchall()
            dconn.close()
            for c, b in rows:
                if c and c.lower() not in seen:
                    d = _damerau(w, c)
                    if d <= 2:
                        seen.add(c.lower())
                        scored.append((d - (b / 1000000.0), c))
        except Exception:
            pass
    scored.sort(key=lambda x: x[0])
    return [c for _, c in scored[:limit]]


@app.post('/api/custom-dict/suggest')
def custom_dict_suggest(body: CustomDictBody):
    word = (body.word or '').strip()
    if not word or _is_phrase(word):
        return {'suggestions': []}
    return {'suggestions': _suggestions(word)}


def _http_get_json(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'KTRT/0.2.0 (local dictionary tool)'})
    try:
        with urllib.request.urlopen(req, timeout=8) as r:
            return json.loads(r.read().decode('utf-8', 'replace'))
    except Exception as e:
        raise RuntimeError(net.describe(e)) from e


@app.post('/api/custom-dict/online')
def custom_dict_online(body: CustomDictBody):
    """在线联想：Datamuse 词汇关系（相关词 / 同音近音 / 形似），按原形查询。

    Wiktionary 已下线：实测经常超时且给出的释义与学习无关，中文机翻改走 /translate。
    """
    word = (body.word or '').strip()
    if not word:
        raise HTTPException(400, '请输入单词')
    canonical = _canonical_word(word)
    q = canonical or word
    out = {'word': word, 'canonical': canonical if canonical.lower() != word.lower() else '',
           'datamuse': {}, 'error': ''}
    try:
        ml = _http_get_json('https://api.datamuse.com/words?ml=' + urllib.parse.quote(q) + '&max=10')
        sl = _http_get_json('https://api.datamuse.com/words?sl=' + urllib.parse.quote(q) + '&max=8')
        sp = _http_get_json('https://api.datamuse.com/words?sp=*' + urllib.parse.quote(q) + '*&max=8')
        out['datamuse'] = {
            'related': [x.get('word') for x in ml[:10] if x.get('word')],
            'sounds_like': [x.get('word') for x in sl[:8] if x.get('word')],
            'spelled_like': [x.get('word') for x in sp[:8] if x.get('word')],
        }
    except Exception as e:
        out['error'] += 'Datamuse 获取失败：%s' % e
    return out


class EnhanceBody(BaseModel):
    word: str = ''
    offline: dict = {}
    online: dict = {}


class TranslateBody(BaseModel):
    text: str = ''
    to: str = ''  # '中文' / '英语'；留空按内容自动判断


@app.post('/api/custom-dict/translate')
def custom_dict_translate(body: TranslateBody):
    """机翻：中文→英语、英语→中文（自动判断方向，也可指定）。"""
    text = (body.text or '').strip()
    if not text:
        raise HTTPException(400, '请输入要翻译的内容')
    if len(text) > 2000:
        raise HTTPException(400, '一次最多翻译 2000 个字符')
    to = (body.to or '').strip()
    if to not in ('中文', '英语'):
        to = '英语' if re.search(r'[\u4e00-\u9fff]', text) else '中文'
    prompt = (
        '把下面的内容翻译成%s，只输出译文本身：不要解释、不要加引号、不要罗列词义、'
        '不要补充原文没有的信息。单词或短语给最常用的对应说法；句子保持原句语气与人称。\n\n%s'
        % (to, text)
    )
    try:
        out = ai.chat([{'role': 'user', 'content': prompt}], max_tokens=1200, temperature=0.2)
    except Exception as e:
        raise HTTPException(502, '翻译失败（需联网且已配置 API Key）：%s' % e)
    return {'text': text, 'to': to, 'translation': (out or '').strip()}


@app.post('/api/custom-dict/enhance')
def custom_dict_enhance(body: EnhanceBody):
    """AI 整理 / 翻译：资料足够时整理排版，资料不足（尤其短语）时用 AI 翻译能力补齐。"""
    word = (body.word or '').strip()
    material = {
        'word': word,
        'offline': body.offline or {},
        'online': body.online or {},
    }
    if _is_phrase(word):
        material['components'] = _phrase_components(word)
    prompt = (
        f'你是英语词典编辑兼翻译。目标词条：「{word}」（可能是单词，也可能是短语）。\n'
        f'原始资料：{json.dumps(material, ensure_ascii=False)[:3500]}\n\n'
        '要求：\n'
        '1. 资料足够时，整理成简洁词卡；资料不足（尤其是短语）时，不要回避：'
        '直接用你的翻译能力给出准确、自然的中文翻译，并标注“AI 翻译”。\n'
        '2. 短语要拆解组成词的含义，说明整体含义、常见用法与语域；可给 1-2 个英文例句及中文翻译。\n'
        '3. 不要编造典故、出处或不存在的事实；不确定的地方标注“不确定”。\n'
        '格式：【释义】/【用法】/【例句】/【同义】/【反义】/【同根】/【组成词】/【备注】，每节一行一条；'
        '短语/短句/习语/俚语可能没有同根词或词形变化，缺失的节直接不写，条目结构与普通单词词条保持一致。'
    )
    try:
        reply = ai.chat([{'role': 'user', 'content': prompt}], max_tokens=900, temperature=0.3)
        return {'ok': True, 'text': reply.strip()}
    except Exception as e:
        return {'ok': False, 'error': 'AI 整理失败（需联网且已配置 API Key）：%s' % e}


def _storm_fetch_online(q):
    """风暴词的在线补充：只取 Datamuse 词汇关系（Wiktionary 已下线：常超时且无中文）。"""
    out = {'datamuse': {}, 'error': ''}
    try:
        ml = _http_get_json('https://api.datamuse.com/words?ml=' + urllib.parse.quote(q) + '&max=12')
        sl = _http_get_json('https://api.datamuse.com/words?sl=' + urllib.parse.quote(q) + '&max=15')
        sp = _http_get_json('https://api.datamuse.com/words?sp=*' + urllib.parse.quote(q) + '*&max=15')
        out['datamuse'] = {
            'related': [x.get('word') for x in ml[:12] if x.get('word')],
            'sounds_like': [x.get('word') for x in sl[:15] if x.get('word')],
            'spelled_like': [x.get('word') for x in sp[:15] if x.get('word')],
        }
    except Exception as e:
        out['error'] += 'Datamuse 获取失败：%s' % e
    return out


def _storm_raw_material(word, q):
    ecdict = {}
    if os.path.exists(db.DICT_DB_PATH):
        try:
            dconn = sqlite3.connect(db.DICT_DB_PATH)
            row = dconn.execute('SELECT * FROM dict WHERE word=? COLLATE NOCASE', (q,)).fetchone()
            dconn.close()
            if row:
                ecdict = {
                    'phonetic': row[1] or '',
                    'translation': row[3] or '',
                    'definition': row[2] or '',
                    'pos': row[4] or '',
                    'exchange': row[5] or '',
                    'exchange_text': _exchange_pretty(row[5] or ''),
                    'oxford': row[7] or '',
                    'collins': row[6] or '',
                }
        except Exception:
            pass
    books = []
    with db.get_conn() as conn:
        for r in conn.execute(
                'SELECT DISTINCT b.name AS book, w.meaning, w.collocations, w.phrases, '
                'w.synonyms, w.antonyms, w.root_words FROM words w '
                'JOIN word_books b ON b.id=w.book_id WHERE w.word=? COLLATE NOCASE', (q,)):
            books.append({k: r[k] or '' for k in r.keys()})
    return {
        'word': word,
        'query': q,
        'ecdict': ecdict,
        'books': books,
        'reference_phrases': _storm_reference_phrases(q),
        'online': _storm_fetch_online(q),
        'confusable_candidates': _suggestions(word, 20),
    }


def _storm_reference_phrases(q):
    """从本地动词短语素材里取与目标词相关的短语/介词搭配（用户上传的《动词短语词典》）。"""
    try:
        with db.get_conn() as conn:
            rows = conn.execute(
                'SELECT phrase, meaning, example FROM reference_phrases '
                'WHERE phrase LIKE ? OR phrase LIKE ? OR phrase LIKE ? LIMIT 20',
                (q + ' %', '% ' + q + ' %', '% ' + q),
            ).fetchall()
        return [{'phrase': r['phrase'], 'meaning': r['meaning'], 'example': r['example']} for r in rows]
    except Exception:
        return []


def _storm_display_word(raw, q):
    """取权威词形：优先词书/ECDICT 里的实际大小写（如 rome→Rome），用于展示与存储。"""
    with db.get_conn() as conn:
        row = conn.execute('SELECT word FROM words WHERE word=? COLLATE NOCASE LIMIT 1', (raw,)).fetchone()
        if row and row['word']:
            return row['word']
    if os.path.exists(db.DICT_DB_PATH):
        try:
            dconn = sqlite3.connect(db.DICT_DB_PATH)
            r = dconn.execute('SELECT word FROM dict WHERE word=? COLLATE NOCASE LIMIT 1', (q,)).fetchone()
            dconn.close()
            if r and r[0] and r[0].lower() == raw.lower():
                return r[0]
        except Exception:
            pass
    return raw


def _storm_markdown(word, c):
    md = ['# ' + word]
    if c.get('phonetic'):
        md.append('音标：' + c.get('phonetic'))
    md.append('')
    if c.get('etymology'):
        md.append('## 百科 · 词源')
        md.append(c.get('etymology'))
        md.append('')
    if c.get('usage'):
        md.append('## 使用场景')
        md.append(c.get('usage'))
        md.append('')
    senses = c.get('senses') or []
    if senses:
        md.append('## 释义')
        for s in senses:
            line = s.get('pos', '') + ' ' + s.get('zh', '')
            if s.get('en'):
                line += '（' + s.get('en') + '）'
            md.append('- ' + line)
        md.append('')
    forms = c.get('forms') or []
    if forms:
        md.append('## 词形变化')
        md += ['- ' + f for f in forms]
        md.append('')
    coll = c.get('collocations') or []
    if coll:
        md.append('## 搭配')
        md += ['- ' + x for x in coll]
        md.append('')
    prep = c.get('prepositions') or []
    if prep:
        md.append('## 介词搭配（词性专属）')
        md += ['- ' + x for x in prep]
        md.append('')
    phr = c.get('phrases') or []
    if phr:
        md.append('## 短语')
        md += ['- ' + x for x in phr]
        md.append('')
    idiom = c.get('idioms') or []
    if idiom:
        md.append('## 俚语 / 习语')
        md += ['- ' + x for x in idiom]
        md.append('')
    family = c.get('word_family') or c.get('root_words') or []
    if family:
        md.append('## 衍生词（各类词性）')
        md += ['- ' + x for x in family]
        md.append('')
    for key, title in (('synonyms', '同义词（含区别）'), ('antonyms', '反义词'), ('near_synonyms', '近义词（含区别）')):
        items = c.get(key) or []
        if items:
            md.append('## ' + title)
            md += ['- ' + x for x in items]
            md.append('')
    conf = c.get('confusables') or []
    if conf:
        md.append('## 形似 / 易混淆词')
        md += ['- ' + x for x in conf]
        md.append('')
    if c.get('sources'):
        md.append('> 来源：' + c.get('sources'))
    return '\n'.join(md).strip() + '\n'


def _storm_list_str(items, drop_codes=False):
    out = []
    for x in items or []:
        if isinstance(x, str):
            s = x.strip()
        elif isinstance(x, dict):
            s = ' '.join(str(x.get(k, '')) for k in ('zh', 'en', 'word', 'definition', 'meaning', 'text') if x.get(k)).strip()
        else:
            continue
        if not s:
            continue
        if drop_codes and re.match(r'^[0-9spdi3]+:', s):
            continue
        out.append(s)
    return out


def _storm_normalize(content):
    """把 AI 可能返回的对象/字符串混杂结果规整成稳定的字符串列表，避免 markdown 拼接崩溃。"""
    for k in ('forms', 'collocations', 'prepositions', 'phrases', 'idioms', 'word_family',
              'synonyms', 'antonyms', 'near_synonyms', 'confusables'):
        content[k] = _storm_list_str(content.get(k), drop_codes=(k == 'forms'))
    senses = content.get('senses')
    if not isinstance(senses, list):
        senses = []
    norm = []
    for s in senses:
        if isinstance(s, dict):
            norm.append({'pos': str(s.get('pos', '') or '').strip(),
                         'zh': str(s.get('zh', '') or '').strip(),
                         'en': str(s.get('en', '') or '').strip()})
        elif isinstance(s, str):
            norm.append({'pos': '', 'zh': s.strip(), 'en': ''})
    content['senses'] = norm
    return content


class StormGenBody(BaseModel):
    word: str = ''
    language: str = '英语'


@app.get('/api/storm')
def list_storm():
    with db.get_conn() as conn:
        rows = conn.execute(
            'SELECT id, word, language, sources, updated_at FROM storm_entries ORDER BY updated_at DESC'
        ).fetchall()
        out = []
        for r in rows:
            pos = _word_positions(conn, r['word'])
            out.append(dict(r) | {
                'in_books': [p['book_name'] for p in pos],
                'positions': pos,
            })
    return out


@app.get('/api/storm/export')
def storm_export(fmt: str = 'md', ids: str = ''):
    if not ids:
        with db.get_conn() as conn:
            rows = conn.execute('SELECT * FROM storm_entries ORDER BY updated_at DESC').fetchall()
    else:
        id_list = [int(x) for x in ids.split(',') if x.strip().isdigit()]
        qs = ','.join('?' for _ in id_list)
        with db.get_conn() as conn:
            rows = conn.execute(f'SELECT * FROM storm_entries WHERE id IN ({qs})', id_list).fetchall() if id_list else []
    if fmt == 'excel':
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '风暴词卡'
        cols = ['单词', '语言', '词源百科', '使用场景', '释义', '词形变化', '搭配', '介词搭配', '短语', '俚语习语', '衍生词', '同义词', '反义词', '近义词', '形似易混淆', '来源', '更新时间']
        ws.append(cols)
        for r in rows:
            c = json.loads(r['content']) if r['content'] else {}
            senses = '；'.join(f"{s.get('pos','')} {s.get('zh','')}" for s in (c.get('senses') or []))
            ws.append([
                r['word'], r['language'], c.get('etymology', ''), c.get('usage', ''), senses,
                '；'.join(c.get('forms') or []),
                '；'.join(c.get('collocations') or []),
                '；'.join(c.get('prepositions') or []),
                '；'.join(c.get('phrases') or []),
                '；'.join(c.get('idioms') or []),
                '；'.join(c.get('word_family') or c.get('root_words') or []),
                '；'.join(c.get('synonyms') or []),
                '；'.join(c.get('antonyms') or []),
                '；'.join(c.get('near_synonyms') or []),
                '；'.join(c.get('confusables') or []),
                r['sources'], r['updated_at'],
            ])
        buf = io.BytesIO()
        wb.save(buf)
        return Response(
            content=buf.getvalue(),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment; filename="storm_export.xlsx"'},
        )
    text = '\n\n---\n\n'.join((r['markdown'] or '') for r in rows)
    return Response(
        content=text,
        media_type='text/markdown; charset=utf-8',
        headers={'Content-Disposition': 'attachment; filename="storm_export.md"'},
    )


@app.get('/api/storm/{sid}')
def get_storm(sid: int):
    with db.get_conn() as conn:
        row = conn.execute('SELECT * FROM storm_entries WHERE id=?', (sid,)).fetchone()
    if row is None:
        raise HTTPException(404, '风暴词卡不存在')
    try:
        content = json.loads(row['content'])
    except Exception:
        content = {}
    return {
        'id': row['id'], 'word': row['word'], 'language': row['language'],
        'content': content, 'markdown': row['markdown'], 'sources': row['sources'],
        'updated_at': row['updated_at'],
    }


@app.post('/api/storm/generate')
def storm_generate(body: StormGenBody):
    raw_word = (body.word or '').strip()
    if not raw_word:
        raise HTTPException(400, '请输入单词')
    language = body.language.strip() or '英语'
    q = _canonical_word(raw_word)
    word = _storm_display_word(raw_word, q)
    raw = _storm_raw_material(word, q)
    prompt = (
        f'你是严谨的英语词典编辑。为单词「{word}」生成一张“风暴词卡”结构化 JSON。'
        f'只依据下面原始资料，不要编造义项或词汇；缺失项用空数组或空字符串。\n'
        f'专有名词要写规范大小写（如 rome 应输出 Rome）。\n'
        f'原始资料：{json.dumps(raw, ensure_ascii=False)[:4000]}\n\n'
        '严格输出一个 JSON 对象，键与要求如下：\n'
        '{"word":"规范大小写后的单词（如 rome→Rome、paris→Paris）",\n'
        '"phonetic":"音标或空字符串",\n'
        '"etymology":"词源与百科背景（1-3 句，中文，无资料则空字符串）",\n'
        '"usage":"使用场景/语域/搭配语境（1-3 句，中文，无则空字符串）",\n'
        '"senses":[{"pos":"词性","zh":"中文释义","en":"英文简释"}],\n'
        '"forms":["词形变化，只写自然语言，如 过去式 abated、复数 phenomena；'
        '绝对不要出现 0:/1:/s: 等原始词形代码"],\n'
        '"collocations":["英文固定搭配 中文翻译"],\n'
        '"prepositions":["该词性专属的介词/小品词搭配 中文翻译，如 abide by 遵守；无则空数组"],\n'
        '"phrases":["普通短语 中文翻译（不要把习语放这里）"],\n'
        '"idioms":["英文习语/俚语 中文翻译（只放真正的习语俚语，无则空数组）"],\n'
        '"word_family":["衍生词 词性. 中文翻译（不同词性，如 Roman n./adj. 罗马人/罗马的）"],\n'
        '"synonyms":["同义词 中文翻译；与靶词的语义区别（每词都写区别）"],\n'
        '"antonyms":["反义词 中文翻译"],\n'
        '"near_synonyms":["近义词 中文翻译；与靶词的语义区别（每词都写区别）"],\n'
        '"confusables":["形似/易混淆词 中文翻译；一句最关键的区分提示"]}\n\n'
        '分类规则：习语/俚语只能进 idioms；普通短语进 phrases；固定搭配进 collocations；'
        '词性专属的介词/小品词搭配进 prepositions；不得互相重复或串类。\n'
        'confusables 由你主动判断“该词需要多少个易混词”：'
        '只保留真正容易混淆的（拼写/读音极近且都是高频常见词，或经典易混对），'
        '普通一字之差但语义无关的词（如 come/home）一律省略；'
        '判断依据是“这个词本身容不容易被拼错/混用”，不要按固定词类/类别套：'
        '无论名词、动词、形容词、副词，还是地名、物品、现象、人物、活动等，'
        '只要不易混就给 0-3 个最易错的即可；确实容易混就给足数量，不设上限；无则空数组。\n'
        '所有同根/同义/反义/近义/衍生词都要给具体中文翻译，近义词和同义词务必写出语义区别。'
    )
    data = None
    last_err = None
    for _ in range(2):
        try:
            reply = ai.chat([{'role': 'user', 'content': prompt}], max_tokens=1800, temperature=0.3)
            data = _extract_json(reply)
            break
        except Exception as e:
            last_err = e
            time.sleep(1.5)
    if data is None:
        raise HTTPException(502, f'AI 生成失败（需联网且已配置 API Key）：{last_err}')
    canonical = str(data.get('word') or '').strip()
    if canonical and re.match(r"^[A-Za-z][A-Za-z\-']*$", canonical):
        word = canonical
    content = dict(data)
    content = _storm_normalize(content)
    content['word'] = word
    content['query'] = q
    content['generated_at'] = time.strftime('%Y-%m-%d %H:%M:%S')
    sources = []
    if raw['ecdict']:
        sources.append('ECDICT')
    if raw['books']:
        sources.append('词书')
    if raw['online']['datamuse'].get('related') or raw['online']['datamuse'].get('sounds_like'):
        sources.append('Datamuse')
    sources.append('AI整理')
    markdown = _storm_markdown(word, content)
    with db._lock:
        with db.get_conn() as conn:
            conn.execute(
                'DELETE FROM storm_entries WHERE lower(word)=lower(?) AND word<>?',
                (word, word),
            )
            conn.execute(
                'INSERT INTO storm_entries(word, language, content, markdown, sources, updated_at) '
                'VALUES(?,?,?,?,?,datetime(\'now\',\'localtime\')) '
                'ON CONFLICT(word, language) DO UPDATE SET content=excluded.content, '
                'markdown=excluded.markdown, sources=excluded.sources, updated_at=excluded.updated_at',
                (word, language, json.dumps(content, ensure_ascii=False), markdown, '、'.join(sources)),
            )
            row = conn.execute('SELECT id FROM storm_entries WHERE word=? AND language=?', (word, language)).fetchone()
    return {'ok': True, 'id': row['id'], 'word': word, 'sources': '、'.join(sources), 'markdown': markdown}


@app.delete('/api/storm/{sid}')
def delete_storm(sid: int):
    with db._lock:
        with db.get_conn() as conn:
            conn.execute('DELETE FROM storm_entries WHERE id=?', (sid,))
    return {'ok': True}


@app.post('/api/custom-dict/lookup')
def custom_dict_lookup(body: CustomDictBody):
    """自定义查词典·速查：离线词典定义 + 词是否在已导入词书 + 是否已收藏。"""
    word = (body.word or '').strip()
    if not word:
        raise HTTPException(400, '请输入单词')
    canonical = _canonical_word(word)
    dict_result = {'found': False, 'phonetic': '', 'translation': '', 'definition': '', 'exchange': ''}
    if os.path.exists(db.DICT_DB_PATH):
        try:
            conn = sqlite3.connect(db.DICT_DB_PATH)
            row = conn.execute('SELECT * FROM dict WHERE word=? COLLATE NOCASE', (word.lower(),)).fetchone()
            conn.close()
            if row:
                dict_result = {
                    'found': True,
                    'phonetic': row[1] or '',
                    'translation': row[3] or '',
                    'definition': row[2] or '',
                    'exchange': row[5] or '',
                    'exchange_text': _exchange_pretty(row[5] or ''),
                }
        except Exception:
            pass
    in_books = []
    favorite = False
    with db.get_conn() as conn:
        for r in conn.execute(
                'SELECT w.id, w.list_no, w.seq, b.id book_id, b.name FROM words w '
                'JOIN word_books b ON b.id=w.book_id WHERE w.word=? COLLATE NOCASE '
                'ORDER BY b.id, w.list_no, w.seq',
                (canonical,)):
            st = conn.execute('SELECT favorite FROM word_status WHERE word_id=?', (r['id'],)).fetchone()
            if st and st['favorite']:
                favorite = True
            in_books.append({'book_id': r['book_id'], 'book_name': r['name'],
                             'list_no': r['list_no'], 'seq': r['seq']})
    components = _phrase_components(word) if _is_phrase(word) else []
    return {
        'word': word,
        'canonical': canonical if canonical.lower() != word.lower() else '',
        'is_phrase': _is_phrase(word),
        'components': components,
        'dict': dict_result,
        'in_books': in_books,
        'favorite': favorite,
        'in_fav_book': any(b['book_name'] == '外部单词收藏册' for b in in_books),
    }


@app.post('/api/custom-dict/favorite')
def custom_dict_favorite(body: CustomDictBody):
    """自定义查词典·收藏：词在任意已导入词书中则标记为收藏。"""
    word = (body.word or '').strip()
    if not word:
        raise HTTPException(400, '请输入单词')
    canonical = _canonical_word(word)
    with db._lock:
        with db.get_conn() as conn:
            row = conn.execute('SELECT id FROM words WHERE word=? COLLATE NOCASE LIMIT 1', (canonical,)).fetchone()
            if row is None:
                raise HTTPException(404, '该词不在任何已导入单词书中，请先「添加」')
            conn.execute('INSERT OR IGNORE INTO word_status(word_id) VALUES(?)', (row['id'],))
            conn.execute('UPDATE word_status SET favorite=1 WHERE word_id=?', (row['id'],))
    return {'ok': True, 'canonical': canonical}


def _insert_external_word(word, phon, meaning, colloc, phras, syns, ants, roots, note_text=''):
    """把词条写入「外部单词收藏册」：每 50 词一个 List；可选把完整 AI 文本存入笔记。"""
    with db._lock:
        with db.get_conn() as conn:
            book = conn.execute("SELECT id FROM word_books WHERE name='外部单词收藏册'").fetchone()
            if book is None:
                cur = conn.execute("INSERT INTO word_books(name, language, source) VALUES('外部单词收藏册','英语','')")
                book_id = cur.lastrowid
            else:
                book_id = book['id']
            row = conn.execute(
                'SELECT list_no, COUNT(*) c FROM words WHERE book_id=? GROUP BY list_no '
                'ORDER BY list_no DESC LIMIT 1', (book_id,)).fetchone()
            if row and row['c'] >= 50:
                list_no, seq = row['list_no'] + 1, 1
            elif row:
                list_no, seq = row['list_no'], row['c'] + 1
            else:
                list_no, seq = 1, 1
            cur = conn.execute(
                'INSERT INTO words(book_id, list_no, seq, word, phonetic, meaning, collocations, '
                'phrases, synonyms, antonyms, root_words) VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                (book_id, list_no, seq, word, phon, meaning, colloc, phras, syns, ants, roots),
            )
            wid = cur.lastrowid
            conn.execute('INSERT OR IGNORE INTO word_status(word_id) VALUES(?)', (wid,))
            if note_text:
                conn.execute(
                    'INSERT INTO word_notes(word_key, content, updated_at) '
                    'VALUES(?,?,datetime(\'now\',\'localtime\')) '
                    'ON CONFLICT(word_key) DO UPDATE SET content=excluded.content, '
                    'updated_at=excluded.updated_at',
                    (_note_key(word), note_text),
                )
    return {'book_name': '外部单词收藏册', 'list_no': list_no, 'seq': seq, 'word_id': wid}


class SaveAiBody(BaseModel):
    word: str = ''
    ai_text: str = ''


def _parse_ai_sections(text):
    def grab(tag):
        m = re.search(r'【' + tag + r'】\s*(.*?)(?=\n?【|$)', text or '', re.S)
        return (m.group(1).strip() if m else '')
    def flat(s):
        return re.sub(r'\s*\n+\s*', '；', s).strip('； ')
    return {
        'meaning': flat(grab('释义')),
        'usage': flat(grab('用法')),
        'examples': flat(grab('例句')),
        'components': flat(grab('组成词')),
        'synonyms': flat(grab('同义')),
        'antonyms': flat(grab('反义')),
        'roots': flat(grab('同根')),
        'remark': flat(grab('备注')),
    }


@app.post('/api/custom-dict/save-ai')
def custom_dict_save_ai(body: SaveAiBody):
    """把 AI 整理/翻译的结果直接存入「外部单词收藏册」，不重复调用 AI。"""
    word = (body.word or '').strip()
    text = (body.ai_text or '').strip()
    if not word or len(word) > 200:
        raise HTTPException(400, '请输入单词或短语')
    if not text:
        raise HTTPException(400, '没有可保存的 AI 释义')
    phrase = _is_phrase(word)
    if phrase:
        if not re.match(r"^[A-Za-z][A-Za-z\-' ,.;:!?]*$", word):
            raise HTTPException(400, '短语/短句格式不支持（仅限英文字母、空格与常用标点）')
        canonical = word
    else:
        if not re.match(r"^[A-Za-z][A-Za-z\-']*$", word):
            raise HTTPException(400, '请输入合法的英文单词')
        canonical = _canonical_word(word)
    with db.get_conn() as conn:
        if conn.execute('SELECT 1 FROM words WHERE word=? COLLATE NOCASE LIMIT 1', (canonical,)).fetchone():
            raise HTTPException(400, '该词已在已导入单词书中，请改用「收藏」')
    sec = _parse_ai_sections(text)
    meaning = sec['meaning'] or text[:300]
    colloc = '；'.join(x for x in (sec['usage'], sec['remark']) if x)
    res = _insert_external_word(
        canonical, '', meaning, colloc, sec['examples'],
        sec['synonyms'], sec['antonyms'], sec['roots'], note_text=text)
    return {'ok': True, 'word': canonical, 'is_phrase': phrase, **res}


@app.post('/api/custom-dict/add')
def custom_dict_add(body: CustomDictBody):
    """自定义查词典·添加：词不在任何已导入词书时，AI 生成词条入默认书。"""
    word = (body.word or '').strip()
    phrase = _is_phrase(word)
    if not word or len(word) > 200:
        raise HTTPException(400, '请输入单词或短语')
    if phrase:
        if not re.match(r"^[A-Za-z][A-Za-z\-' ,.;:!?]*$", word):
            raise HTTPException(400, '短语/短句格式不支持（仅限英文字母、空格与常用标点）')
        canonical = word
    else:
        if not re.match(r"^[A-Za-z][A-Za-z\-']*$", word):
            raise HTTPException(400, '请输入合法的英文单词')
        canonical = _canonical_word(word)
    with db.get_conn() as conn:
        exists = conn.execute('SELECT 1 FROM words WHERE word=? COLLATE NOCASE LIMIT 1', (canonical,)).fetchone()
    if exists:
        raise HTTPException(400, '该词已在已导入单词书中，请改用「收藏」')
    word = canonical
    if phrase:
        comps = '；'.join(f"{c['word']} {c['translation'] or '（本地未收录）'}" for c in _phrase_components(word))
        prompt = (
            f'英语短语/短句/习语/俚语「{word}」：只输出一个 JSON 对象，不要任何其他文字：'
            '{"type": "类型，只能是 俚语 / 习语 / 固定句式 / 表达 / 短语 之一", '
            '"translation": "整体中文翻译（1-3 个义项用分号分隔）", '
            '"usage": "使用场景与语域说明（中文 1-2 句；不确定则空字符串）", '
            '"breakdown": "拆解：逐词或结构的含义、为什么整体是这个意思（中文，用分号分隔；无则空字符串）"}。'
            f'参考组成词释义：{comps}。'
            '要求：翻译准确自然；不确定的内容宁可留空，不要编造典故、出处或事实。'
        )
    else:
        prompt = (
            f'为英语单词「{word}」生成词条数据，只输出一个 JSON 对象，不要任何其他文字：'
            '{"phonetic": "国际音标", "meaning": "词性. 中文释义", '
            '"collocations": "两个常见搭配（英文短语，用分号分隔）", '
            '"phrases": "两个常见短语（英文短语，用分号分隔）", '
            '"synonyms": "2-3 个同义词（英文，用分号分隔）", '
            '"antonyms": "1-2 个反义词（英文，用分号分隔）", '
            '"root_words": "2-3 个同根词（英文，用分号分隔）"}。'
            '要求：字段齐全、内容真实准确、不要编造。'
        )
    try:
        raw = ai.chat([{'role': 'user', 'content': prompt}], max_tokens=800, temperature=0.3)
        data = _extract_json(raw)
    except Exception as e:
        raise HTTPException(502, f'AI 生成失败（需联网且已配置 API Key）：{e}')
    if phrase:
        # 词组/句子只保留释义：类型 + 译文 + 场景 + 拆解，其余列留空
        lines = []
        tag = str(data.get('type') or '').strip()
        if tag:
            lines.append('【%s】' % tag)
        for label, key in (('译文', 'translation'), ('场景', 'usage'), ('拆解', 'breakdown')):
            v = str(data.get(key) or '').strip()
            if v:
                lines.append('%s：%s' % (label, v))
        meaning = '\n'.join(lines)
        phon = colloc = phras = syns = ants = roots = ''
    else:
        phon = str(data.get('phonetic') or '').strip()
        meaning = str(data.get('meaning') or '').strip()
        colloc = str(data.get('collocations') or '').strip()
        phras = str(data.get('phrases') or '').strip()
        syns = str(data.get('synonyms') or '').strip()
        ants = str(data.get('antonyms') or '').strip()
        roots = str(data.get('root_words') or '').strip()
    res = _insert_external_word(word, phon, meaning, colloc, phras, syns, ants, roots)
    return {'ok': True, 'word': word, 'is_phrase': phrase, **res}


@app.get('/api/references')
def references(word: str = Query(''), limit: int = Query(10), keys: str = Query('')):
    """动词短语参考：按短语首词（动词）匹配，支持原形归一。

    有内置短语库（含中文释义/例句中译/语域/学习价值）时返回分义项结构，
    否则回退到旧的英文参考表。
    """
    wl = word.strip().lower()
    limit = max(1, min(limit, 20))
    if phrasal.BY_KEY:
        # 优先用导入时写好的命中结果（keys），没有再做实时检索
        stored = [k.strip() for k in (keys or '').split(';') if k.strip()]
        if stored:
            return phrasal.entries_for_keys(stored, limit)
        canonical = wl if ' ' in wl else (_canonical_word(wl) or wl).lower()
        hits = list(phrasal.BY_HEAD.get(wl, []))
        if canonical != wl:
            hits += phrasal.BY_HEAD.get(canonical, [])
        seen, out = set(), []
        for g in hits:
            if g['key'] in seen:
                continue
            seen.add(g['key'])
            out.append(g)
        return out[:limit]
    with db.get_conn() as conn:
        if wl:
            canonical = wl if ' ' in wl else (_canonical_word(wl) or wl).lower()
            rows = conn.execute(
                "SELECT id, phrase, meaning, example, source FROM reference_phrases "
                "WHERE lower(substr(phrase, 1, instr(phrase || ' ', ' ') - 1)) IN (?, ?) "
                "ORDER BY length(phrase), phrase, id LIMIT ?",
                (wl, canonical, limit),
            ).fetchall()
        else:
            rows = conn.execute(
                'SELECT id, phrase, meaning, example, source FROM reference_phrases '
                'ORDER BY phrase, id LIMIT ?',
                (limit,),
            ).fetchall()
    return [dict(r) for r in rows]


@app.get('/api/tts')
def read_aloud(text: str = Query(''), lang: str = Query('英语')):
    if not text.strip():
        raise HTTPException(400, '没有可朗读的文本')
    try:
        voice_key = db.get_setting('tts_voice_fr' if lang == '法语' else 'tts_voice_en',
                                   '女声' if lang == '法语' else '美音·男')
        rate = db.get_setting('tts_rate', '0')
        pitch = db.get_setting('tts_pitch', '0')
        volume = db.get_setting('tts_volume', '100')
        path = tts.synthesize(text.strip(), lang, voice_key, rate, pitch, volume)
    except Exception as e:
        raise HTTPException(502, f'语音合成失败：{e}')
    return FileResponse(path, media_type='audio/mpeg', headers={'Cache-Control': 'no-store'})


def _get_settings():
    cfg = ai.current_config()
    return {
        # 不回传明文 Key：本机任意进程、任意网页都不该读到它。
        # 前端只需要知道"有没有配过"，改 Key 时再整串覆盖。
        'api_key': '',
        'api_key_set': bool(cfg['api_key']),
        'base_url': cfg['base_url'],
        'model': cfg['model'],
        'vendor': cfg['vendor'],
        'tts_provider': db.get_setting('tts_provider', 'edge-tts'),
        'tts_voice_en': db.get_setting('tts_voice_en', '美音·男'),
        'tts_voice_fr': db.get_setting('tts_voice_fr', '女声'),
        'tts_rate': db.get_setting('tts_rate', '0'),
        'tts_pitch': db.get_setting('tts_pitch', '0'),
        'tts_volume': db.get_setting('tts_volume', '100'),
        'theme': db.get_setting('theme', 'dark-blue'),
        'theme_page': db.get_setting('theme_page', 'normal'),
        'auto_update_check': db.get_setting('auto_update_check', '1'),
        'update_snooze': db.get_setting('update_snooze', ''),
    }


class SettingsBody(BaseModel):
    api_key: str = ''
    clear_api_key: bool = False
    base_url: str = ''
    model: str = ''
    vendor: str = 'ds'
    tts_provider: str = 'edge-tts'
    tts_voice_en: str = '美音·男'
    tts_voice_fr: str = '女声'
    tts_rate: str = '0'
    tts_pitch: str = '0'
    tts_volume: str = '100'
    theme: str = 'dark-blue'
    theme_page: str = 'normal'


@app.get('/api/settings')
def get_settings():
    return _get_settings()


@app.post('/api/settings')
def save_settings(body: SettingsBody):
    key = body.api_key.strip()
    if body.clear_api_key:
        db.set_setting('api_key', '')
    elif key:
        # 留空表示"不修改"：前端拿不到明文，空值不能当成删除 Key
        db.set_setting('api_key', key)
    db.set_setting('base_url', body.base_url.strip().rstrip('/'))
    db.set_setting('model', body.model.strip())
    db.set_setting('vendor', body.vendor.strip() or 'ds')
    db.set_setting('tts_provider', body.tts_provider.strip() or 'edge-tts')
    db.set_setting('tts_voice_en', body.tts_voice_en.strip() or '美音·男')
    db.set_setting('tts_voice_fr', body.tts_voice_fr.strip() or '女声')
    db.set_setting('tts_rate', body.tts_rate.strip() or '0')
    db.set_setting('tts_pitch', body.tts_pitch.strip() or '0')
    db.set_setting('tts_volume', body.tts_volume.strip() or '100')
    db.set_setting('theme', body.theme.strip() or 'dark-blue')
    db.set_setting('theme_page', body.theme_page.strip() or 'normal')
    return _get_settings()


@app.get('/api/update/status')
def update_status():
    """检查更新：当前版本 / 最新 Release / 可用的补丁与整包 / 已装补丁。"""
    return updater.status()


class UpdateApplyBody(BaseModel):
    kind: str = 'patch'


@app.post('/api/update/apply')
def update_apply(body: UpdateApplyBody):
    """一键更新：patch 走热补丁（不用重启），full 下载安装包静默安装并自动重开。"""
    kind = (body.kind or '').strip()
    if kind not in ('patch', 'full'):
        raise HTTPException(400, '未知的更新类型')
    info = updater.status(deep=False)
    latest = info.get('latest') or {}
    target = info.get('patch') if kind == 'patch' else info.get('installer')
    if not target:
        raise HTTPException(400, '没有可用的更新内容，请先点「检查更新」')
    if kind == 'full' and info.get('mode') != 'packaged':
        raise HTTPException(400, '源码版请用 git pull 更新，一键安装只对安装版可用')
    try:
        return updater.start(kind, target['url'], latest.get('version') or info['current_version'])
    except Exception as e:
        raise HTTPException(400, str(e))


@app.get('/api/update/progress')
def update_progress():
    out = updater.job()
    out['overlay'] = updater.overlay_active()
    return out


class UpdatePrefsBody(BaseModel):
    auto_check: str = ''
    snooze: str = ''


@app.post('/api/update/prefs')
def update_prefs(body: UpdatePrefsBody):
    """更新页偏好：启动时自动检查、本版本不再提醒。"""
    if body.auto_check:
        db.set_setting('auto_update_check', '1' if body.auto_check == '1' else '0')
    if body.snooze:
        db.set_setting('update_snooze', body.snooze.strip())
    return {'auto_update_check': db.get_setting('auto_update_check', '1'),
            'update_snooze': db.get_setting('update_snooze', '')}


@app.post('/api/ai/test')
def ai_test():
    try:
        reply = ai.chat([{'role': 'user', 'content': '只回复两个汉字：成功'}], max_tokens=10)
        return {'ok': True, 'reply': reply.strip()}
    except Exception as e:
        return {'ok': False, 'error': str(e)}


if __name__ == '__main__':
    import uvicorn
    uvicorn.run(app, host='127.0.0.1', port=8000)
