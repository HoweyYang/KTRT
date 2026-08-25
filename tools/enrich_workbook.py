#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
KTRT 词库补全工具
--------------------------------
1. 清理词库里的水印广告（如「需要完整的词汇表格加微信：ｃ ２００２１９５９」）。
2. 按 GRE 格式为「搭配」「短语」逐条补中文翻译（English 中文；English 中文）。
3. 可选把补全结果回写到 data/ktrt.db（按 单词书+单词 匹配，不重复导入）。

用法示例：
    python tools/enrich_workbook.py --file wordbooks/IELTS_Wordbook.xlsx --book 雅思词汇真经 --update-db
    python tools/enrich_workbook.py --file wordbooks/KAOYAN_Wordbook.xlsx --book 考研英语词汇词根+联想记忆法 --update-db
    python tools/enrich_workbook.py --file xxx.xlsx --book 外部单词收藏册 --update-db

说明：
- 已含中文的条目原样保留，不会被重复翻译。
- 翻译结果缓存在 data/enrich_cache.json，按短语文本去重，中断后重跑可续传。
- 翻译调用当前「设置」里配置的 AI（默认 DeepSeek），Key 在 data/config.json。
"""

import argparse
import json
import os
import re
import sys
import time

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from backend import ai  # noqa: E402

DB_PATH = os.path.join(ROOT, 'data', 'ktrt.db')
CACHE_PATH = os.path.join(ROOT, 'data', 'enrich_cache.json')

CJK_RE = re.compile(r'[\u4e00-\u9fff]')
# 水印：需要……加微信：ｃ ２００２１９５９（全/半角数字、空格均可）
WATERMARK_RE = re.compile(r'\s*需要[^加]{0,40}加微信[:：]\s*[ｃcCＣ]\s*[０-９0-9\s]+')
# 兼容旧水印「加微信：xxx」其它变体（如只留数字）
WATERMARK_TAIL_RE = re.compile(r'\s*加微信[:：]\s*[ｃcCＣ]?\s*[０-９0-9\s]+$')

BATCH = 80
MAX_TOKENS = 2048
MAX_RETRY = 3


def has_cjk(s):
    return bool(CJK_RE.search(s))


def split_items(v):
    if not v:
        return []
    return [x.strip() for x in str(v).split('；') if x.strip()]


def clean_cell(v):
    if not v:
        return v
    s = str(v)
    s2 = WATERMARK_RE.sub('', s)
    s2 = WATERMARK_TAIL_RE.sub('', s2)
    if s2 != s:
        s2 = s2.rstrip(' \t;；。.')
    return s2 if s2 != s else v


def load_cache():
    if os.path.exists(CACHE_PATH):
        try:
            with open(CACHE_PATH, encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def save_cache(cache):
    tmp = CACHE_PATH + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(cache, f, ensure_ascii=False, indent=1)
    os.replace(tmp, CACHE_PATH)


def translate_batch(items):
    """返回 {原短语: 中文翻译}。只翻译纯英文条目。"""
    lines = '\n'.join(f'{i}: {it}' for i, it in enumerate(items))
    messages = [
        {'role': 'system',
         'content': '你是英语词汇编辑，精通中英翻译。只输出 JSON，不要任何解释或多余文字。'},
        {'role': 'user',
         'content': '把下面每条英文搭配/短语翻译成简洁自然的中文。\n'
                    '输出 JSON 对象，键为序号，值为对应中文翻译。\n\n' + lines},
    ]
    last_err = None
    for attempt in range(MAX_RETRY):
        try:
            txt = ai.chat(messages, max_tokens=MAX_TOKENS, temperature=0.3).strip()
            txt = re.sub(r'^```[a-zA-Z]*\s*', '', txt)
            txt = re.sub(r'\s*```$', '', txt)
            data = json.loads(txt)
            result = {}
            for k, v in data.items():
                try:
                    result[items[int(k)]] = str(v).strip()
                except Exception:
                    continue
            if result:
                return result
        except Exception as e:  # noqa: BLE001
            last_err = e
            time.sleep(2 + attempt * 3)
    print(f'    [warn] 批次翻译失败，{len(items)} 条留待下次重跑: {last_err}')
    return {}


def collect_tasks(ws, cache):
    """返回 tasks: [(row, col, item)] 和 unique_items: [item...]"""
    tasks = []
    seen = set()
    for row in ws.iter_rows(min_row=2):
        if not row[0].value:
            continue
        for col in (3, 4):  # 搭配 / 短语
            for item in split_items(row[col].value):
                if has_cjk(item):
                    continue
                if item in cache:
                    continue
                tasks.append((row, col, item))
                if item not in seen:
                    seen.add(item)
    return tasks, list(seen)


def translate_all(tasks, unique_items):
    cache = load_cache()
    done_new = 0
    for start in range(0, len(unique_items), BATCH):
        chunk = unique_items[start:start + BATCH]
        result = translate_batch(chunk)
        for item, zh in result.items():
            if zh and not has_cjk(item + ' ' + zh):
                zh = item  # 兜底：模型没给出中文就保留原文
            cache[item] = f'{item} {zh}' if zh and zh != item else item
            done_new += 1
        save_cache(cache)
        done = start + len(chunk)
        print(f'  [{done}/{len(unique_items)}] 本批翻译 {len(result)} 条，'
              f'累计新增 {done_new} 条，缓存共 {len(cache)} 条')
    return cache


def apply_to_sheet(ws, cache):
    changed = 0
    for row in ws.iter_rows(min_row=2):
        if not row[0].value:
            continue
        for col in (3, 4):
            parts = split_items(row[col].value)
            new_parts = []
            for p in parts:
                if not has_cjk(p) and p in cache:
                    new_parts.append(cache[p])
                    changed += 1
                else:
                    new_parts.append(p)
            if new_parts != parts:
                row[col].value = '；'.join(new_parts)
    return changed


def update_db(book_name, ws):
    import sqlite3
    con = sqlite3.connect(DB_PATH)
    cur = con.cursor()
    cur.execute('SELECT id FROM word_books WHERE name=?', (book_name,))
    row = cur.fetchone()
    if not row:
        print(f'[db] 找不到词书「{book_name}」，跳过数据库回写')
        con.close()
        return
    book_id = row[0]
    updated = 0
    for ws_row in ws.iter_rows(min_row=2):
        word = ws_row[0].value
        if not word:
            continue
        coll = ws_row[3].value
        phr = ws_row[4].value
        meaning = ws_row[2].value
        cur.execute(
            'UPDATE words SET meaning=?, collocations=?, phrases=? '
            'WHERE book_id=? AND word=?',
            (meaning, coll, phr, book_id, str(word).strip()),
        )
        updated += cur.rowcount
    con.commit()
    con.close()
    print(f'[db] 「{book_name}」回写完成，更新 {updated} 行')


def main():
    ap = argparse.ArgumentParser(description='KTRT 词库补全（水印清理 + 搭配/短语中文翻译 + 数据库回写）')
    ap.add_argument('--file', required=True, help='要处理的 xlsx 路径')
    ap.add_argument('--book', default='', help='词书名（与文件内【单词书】一致），配合 --update-db 使用')
    ap.add_argument('--update-db', action='store_true', help='处理后回写 data/ktrt.db')
    ap.add_argument('--skip-translate', action='store_true', help='只清理水印，不翻译')
    args = ap.parse_args()

    path = os.path.abspath(args.file)
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    print(f'处理: {path}（{ws.title}，{ws.max_row} 行）')

    # 1. 水印清理（全表扫）
    wm_rows = 0
    for row in ws.iter_rows():
        for c in row:
            cleaned = clean_cell(c.value)
            if cleaned != c.value:
                c.value = cleaned
                wm_rows += 1
    print(f'[watermark] 清理 {wm_rows} 个单元格')

    # 2. 翻译
    if not args.skip_translate:
        tasks, unique_items = collect_tasks(ws, load_cache())
        print(f'[translate] 待翻译条目 {len(tasks)} 条（去重后 {len(unique_items)} 条）')
        if unique_items:
            cache = translate_all(tasks, unique_items)
        else:
            cache = load_cache()
        changed = apply_to_sheet(ws, cache)
        print(f'[translate] 已补中文 {changed} 处')
    else:
        cache = load_cache()
        changed = apply_to_sheet(ws, cache)
        print(f'[translate] 已按缓存补中文 {changed} 处')

    wb.save(path)
    print(f'[save] 已保存 {path}')

    if args.update_db:
        book = args.book or ws.cell(row=2, column=10).value
        update_db(book, ws)

    # 3. 复核：还剩多少未翻译条目 / 水印
    leftover = 0
    wm_left = 0
    for row in ws.iter_rows(min_row=2):
        if not row[0].value:
            continue
        for col in (2, 3, 4):
            v = str(row[col].value or '')
            if '微信' in v:
                wm_left += 1
        for col in (3, 4):
            for item in split_items(row[col].value):
                if not has_cjk(item):
                    leftover += 1
    print(f'[check] 剩余水印单元格 {wm_left}，剩余未翻译条目 {leftover}')


if __name__ == '__main__':
    main()
